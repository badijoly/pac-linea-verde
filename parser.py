# -*- coding: utf-8 -*-
"""
Módulo de extracción y cálculo de métricas para Participantes PAC 2026.
Implementa normalización canónica de establecimientos y evaluación
cualitativa rigurosa según directrices pedagógicas.
"""

import openpyxl
import unicodedata
import os
from collections import defaultdict, Counter

def clean_str(s):
    if s is None:
        return ''
    return str(s).strip()

def strip_accents(s):
    if not s:
        return ''
    return ''.join(c for c in unicodedata.normalize('NFD', str(s)) if unicodedata.category(c) != 'Mn').strip().lower()

LINEA_VERDE_SCHOOLS = {
    'Hito 3': [
        'Escuela La Bandera',
        'Escuela Presbiteriana',
        'Escuela Héroes de la Concepción',
        'Escuela Rómulo Peña'
    ],
    'Hito 4': [
        'Liceo Bicentenario Andrés Sabella',
        'Greenhill School',
        'Colegio Bet-el'
    ],
    'Global': [
        'Escuela La Bandera',
        'Escuela Presbiteriana',
        'Escuela Héroes de la Concepción',
        'Escuela Rómulo Peña',
        'Liceo Bicentenario Andrés Sabella',
        'Greenhill School',
        'Colegio Bet-el'
    ]
}
ALL_LINEA_VERDE = set(LINEA_VERDE_SCHOOLS['Global'])

def canonical_school_name(name):
    if not name:
        return ''
    n = str(name).strip().replace('"', '').replace("'", "")
    nc = strip_accents(n)
    if 'heroes' in nc:
        return 'Escuela Héroes de la Concepción'
    if 'sabella' in nc:
        return 'Liceo Bicentenario Andrés Sabella'
    if 'bet-el' in nc or 'bet el' in nc or 'betel' in nc:
        return 'Colegio Bet-el'
    if 'bandera' in nc:
        return 'Escuela La Bandera'
    if 'presbiteriana' in nc:
        return 'Escuela Presbiteriana'
    if 'romulo' in nc:
        return 'Escuela Rómulo Peña'
    if 'greenhill' in nc:
        return 'Greenhill School'
    return n

FEMALE_NAMES = {
    'amanda', 'danna', 'isabella', 'paulina', 'sophia', 'xochtiel', 'abigail', 'ayleen', 
    'estefania', 'eysa', 'fabiola', 'gladys', 'mariana', 'matilde', 'matilda', 'monserrat',
    'pamela', 'valentina', 'dafne', 'lia', 'leah', 'julieta', 'kamila', 'janely', 'isidora',
    'salome', 'sara', 'daniela', 'fabrianny', 'fernanda', 'susy', 'monica', 'carol', 'pia',
    'guisselle', 'jocelyn', 'rosmery', 'javiera', 'mariel', 'kathleen', 'martina', 'ignacia',
    'constanza', 'catalina', 'camila', 'antonia', 'francisca', 'florencia', 'lucia', 'valeria',
    'sofia', 'emilia', 'trinidad', 'maite', 'mia', 'zoe', 'antonella', 'dominic', 'breymarw',
    'gioberlys', 'nahuel'
}

MALE_NAMES = {
    'agustin', 'axel', 'deiby', 'ignacio', 'raul', 'benjamin', 'demian',
    'matias', 'maximiliano', 'pablo', 'emilio', 'tomas', 'gian',
    'gabriel', 'carlos', 'mateo', 'lucas', 'joaquin', 'vicente', 'martin', 'santiago',
    'alonso', 'cristobal', 'bruno', 'dilan', 'ian', 'alvaro', 'rodrigo', 'pedro',
    'nicolas', 'sebastian', 'felipe', 'diego', 'claudio'
}

def infer_gender(full_name):
    first = strip_accents(full_name).split()[0] if full_name else ''
    if first in FEMALE_NAMES:
        return 'Femenino'
    if first in MALE_NAMES:
        return 'Masculino'
    if first.endswith('a'):
        return 'Femenino'
    if first.endswith(('o', 'el', 'an', 'or', 'in', 'en')):
        return 'Masculino'
    return 'Femenino' if first.endswith(('i', 'y')) else 'Masculino'

KEYS = {
    'Hito 3': ['d', 'a', 'b', 'c', 'c'],
    'Hito 4': ['b', 'c', 'a', 'b', 'b']
}

NON_INFORMATIVE = [
    'no se', 'nada', 'ninguna', 'no me acuerdo', 'nose', 'no responde', 
    'no se como decirlo', 'no se como puedo decirlo', 'no se explicarlo', 
    'no se que poner', 'no sabria decir', 'no responder', 'ninguno', 'no lo se'
]

def is_non_informative(text):
    if not text:
        return True
    t = strip_accents(str(text)).strip(' .,;:-_')
    if len(t) < 2:
        return True
    for b in NON_INFORMATIVE:
        if t == b or t.startswith(b + ' ') or t.endswith(' ' + b):
            return True
    return False

def evaluate_conceptual_level(text, hito):
    if is_non_informative(text):
        return 0, 'No sabe / Sin respuesta'
    tc = strip_accents(text)
    kw_h3 = ['ecosistema', 'biodiversidad', 'gaviotin', '50 gr', '50 gramos', 'microorganismo', 'flora', 'fauna', 'salina', 'desierto costero', 'cadena trofica', 'humedal', 'reserva', 'portada', 'endemica', 'migratoria', 'piqueros', 'piquero', 'pelicano']
    kw_h4 = ['atmosfera', 'escudo', 'capa de ozono', 'gases', 'efecto invernadero', 'radiacion', 'uv', 'contaminacion luminica', 'astronomia', 'telescopio', 'paranal', 'luz azul', 'cielos limpios', 'patrimonio', 'postes de luz mirando hacia arriba', 'postes']
    kw = kw_h3 if hito == 'Hito 3' else kw_h4
    matches = sum(1 for k in kw if k in tc)
    words = len(tc.split())
    if matches >= 1 and (words >= 4 or 'gaviotin' in tc or '50' in tc or 'ecosistema' in tc or 'escudo' in tc or 'postes' in tc):
        return 3, 'Apropiación Científica / Específica'
    elif matches >= 1 or words >= 3:
        return 2, 'Comprensión Intermedia / Intuitiva'
    else:
        return 1, 'Opinión Básica / Opinión General'

def compare_open_answers(pre_text, post_text, hito):
    pre_non = is_non_informative(pre_text)
    post_non = is_non_informative(post_text)
    
    # 1. Ambos sin información o "no sé"
    if pre_non and post_non:
        return {
            'status': 'Sin Evidencia',
            'detail': 'Sin respuesta en PRE ni POST (Requiere revisión con profesor/a líder)',
            'requires_review': True,
            'pre_level': 0,
            'post_level': 0
        }
        
    # 2. Inicialmente dijo "no sé" o en blanco, y en post sí tiene respuesta real
    if pre_non and not post_non:
        lvl, lbl = evaluate_conceptual_level(post_text, hito)
        return {
            'status': 'Mejoró',
            'detail': f'Avance conceptual real: Pasa de no responder a expresar {lbl}',
            'requires_review': False,
            'pre_level': 0,
            'post_level': lvl
        }
        
    # 3. Tenía respuesta y en post dejó de responder o puso "no sé"
    if not pre_non and post_non:
        lvl_pre, lbl_pre = evaluate_conceptual_level(pre_text, hito)
        return {
            'status': 'Retrocedió',
            'detail': f'Dejó de responder en el POST (Basal era {lbl_pre})',
            'requires_review': True,
            'pre_level': lvl_pre,
            'post_level': 0
        }
        
    # 4. Ambos tienen respuesta con contenido
    lvl_pre, lbl_pre = evaluate_conceptual_level(pre_text, hito)
    lvl_post, lbl_post = evaluate_conceptual_level(post_text, hito)
    
    if lvl_post > lvl_pre:
        return {
            'status': 'Mejoró',
            'detail': f'Mayor precisión y vocabulario ({lbl_pre} ➔ {lbl_post})',
            'requires_review': False,
            'pre_level': lvl_pre,
            'post_level': lvl_post
        }
    elif lvl_post == lvl_pre:
        return {
            'status': 'Mantuvo',
            'detail': f'Mantuvo respuesta informada ({lbl_post})',
            'requires_review': False,
            'pre_level': lvl_pre,
            'post_level': lvl_post
        }
    else:
        return {
            'status': 'Retrocedió',
            'detail': f'Menor nivel de detalle ({lbl_pre} ➔ {lbl_post})',
            'requires_review': False,
            'pre_level': lvl_pre,
            'post_level': lvl_post
        }

def parse_excel_data(file_path):
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Archivo no encontrado en {file_path}")
        
    wb = openpyxl.load_workbook(file_path, data_only=True)
    
    # 1. Metadatos de Colegios y Líderes con normalización canónica
    schools = {}
    if 'Participantes 2026' in wb.sheetnames:
        s_part = wb['Participantes 2026']
        for r in range(2, s_part.max_row + 1):
            raw_sch = clean_str(s_part.cell(r, 2).value)
            if not raw_sch:
                continue
            sch = canonical_school_name(raw_sch)
            if sch not in ALL_LINEA_VERDE:
                continue
            if sch not in schools:
                schools[sch] = {
                    'nombre': sch,
                    'equipo': s_part.cell(r, 1).value,
                    'direccion': clean_str(s_part.cell(r, 3).value),
                    'dependencia': clean_str(s_part.cell(r, 4).value) or 'Municipal',
                    'certificacion': clean_str(s_part.cell(r, 5).value) or 'No',
                    'participacion_previa': clean_str(s_part.cell(r, 6).value) or 'No',
                    'lider1': clean_str(s_part.cell(r, 8).value),
                    'email1': clean_str(s_part.cell(r, 9).value),
                    'tel1': clean_str(s_part.cell(r, 10).value),
                    'lider2': clean_str(s_part.cell(r, 11).value),
                    'email2': clean_str(s_part.cell(r, 12).value),
                    'tel2': clean_str(s_part.cell(r, 13).value),
                }

    # Cruce con Resumen Líderes
    if 'Resumen Líderes' in wb.sheetnames:
        s_lid = wb['Resumen Líderes']
        for r in range(2, s_lid.max_row + 1):
            raw_sch = clean_str(s_lid.cell(r, 2).value)
            if not raw_sch:
                continue
            sch = canonical_school_name(raw_sch)
            if sch not in ALL_LINEA_VERDE:
                continue
            if sch not in schools:
                schools[sch] = {
                    'nombre': sch,
                    'dependencia': 'Municipal' if 'Escuela' in sch or 'Liceo' in sch else 'Particular Subvencionado',
                    'certificacion': 'No',
                    'lider1': '', 'email1': '', 'tel1': '',
                    'lider2': '', 'email2': '', 'tel2': ''
                }
            # Asignar líderes desde la pestaña Resumen Líderes si están disponibles
            l1 = clean_str(s_lid.cell(r, 3).value)
            t1 = clean_str(s_lid.cell(r, 4).value)
            e1 = clean_str(s_lid.cell(r, 5).value)
            l2 = clean_str(s_lid.cell(r, 6).value)
            t2 = clean_str(s_lid.cell(r, 7).value)
            e2 = clean_str(s_lid.cell(r, 8).value)
            
            if l1:
                schools[sch]['lider1'] = l1
                schools[sch]['tel1'] = t1
                schools[sch]['email1'] = e1
            if l2:
                schools[sch]['lider2'] = l2
                schools[sch]['tel2'] = t2
                schools[sch]['email2'] = e2

    hitos_data = {}
    all_participants = {}

    hitos_config = [
        ('Hito 3', 'Encuestas Hito 3_Verde', 'Los otros vecinos de Antofagasta'),
        ('Hito 4', 'Encuestas Hito 4_Verde', 'Proteger nuestro escudo invisible')
    ]

    for hito_name, sheet_name, desafio_name in hitos_config:
        if sheet_name not in wb.sheetnames:
            continue
            
        ws = wb[sheet_name]
        key = KEYS[hito_name]
        q_titles = [clean_str(ws.cell(1, c).value) for c in range(5, 10)]
        open_q_titles = [clean_str(ws.cell(1, c).value) for c in range(10, 14)]
        
        p_dict = defaultdict(dict)
        for r in range(2, ws.max_row + 1):
            name = clean_str(ws.cell(r, 1).value)
            if not name:
                continue
            role = clean_str(ws.cell(r, 2).value) or 'Estudiante'
            school = canonical_school_name(clean_str(ws.cell(r, 3).value))
            if school not in ALL_LINEA_VERDE:
                continue
                
            pp = clean_str(ws.cell(r, 4).value).upper()
            if pp not in ['PRE', 'POST']:
                continue
            
            mc_ans = [clean_str(ws.cell(r, c).value).lower() for c in range(5, 10)]
            open_ans = [clean_str(ws.cell(r, c).value) for c in range(10, 14)]
            
            correct = [1 if mc_ans[i] == key[i] else 0 for i in range(5)]
            score = sum(correct)
            
            p_dict[name]['name'] = name
            p_dict[name]['role'] = role
            if school:
                p_dict[name]['school'] = school
            p_dict[name][pp] = {
                'row': r,
                'mc': mc_ans,
                'correct': correct,
                'score': score,
                'open': open_ans
            }
            
        paired_list = []
        for name, d in p_dict.items():
            gender = infer_gender(name)
            sch = d.get('school', 'Sin Asignar')
            
            if name not in all_participants:
                all_participants[name] = {
                    'name': name,
                    'role': d.get('role', 'Estudiante'),
                    'school': sch,
                    'gender': gender,
                    'hitos': {}
                }
            
            has_pre = 'PRE' in d
            has_post = 'POST' in d
            
            if has_pre and has_post:
                pre_score = d['PRE']['score']
                post_score = d['POST']['score']
                delta = post_score - pre_score
                mc_status = 'Mejoró' if delta > 0 else ('Mantuvo' if delta == 0 else 'Retrocedió')
                
                # Evolución por pregunta de alternativas
                q_evo = []
                for i in range(5):
                    cp = d['PRE']['correct'][i]
                    cpo = d['POST']['correct'][i]
                    if cp == 0 and cpo == 1:
                        q_evo.append('Mejoró')
                    elif cp == 1 and cpo == 1:
                        q_evo.append('Mantuvo Correcto')
                    elif cp == 0 and cpo == 0:
                        q_evo.append('Mantuvo Incorrecto')
                    else:
                        q_evo.append('Retrocedió')
                        
                # Análisis cualitativo riguroso de preguntas abiertas
                open_analysis = []
                qual_stats = {'Mejoró': 0, 'Mantuvo': 0, 'Retrocedió': 0, 'Sin Evidencia': 0}
                
                for i in range(4):
                    t_pre = d['PRE']['open'][i]
                    t_post = d['POST']['open'][i]
                    res = compare_open_answers(t_pre, t_post, hito_name)
                    qual_stats[res['status']] += 1
                    
                    open_analysis.append({
                        'question': open_q_titles[i],
                        'pre_text': t_pre,
                        'post_text': t_post,
                        'status': res['status'],
                        'detail': res['detail'],
                        'requires_review': res['requires_review']
                    })
                
                # Determinación del Estado Real de Aprendizajes (Integrado)
                if qual_stats['Sin Evidencia'] == 4:
                    real_status = 'Sin evidencia suficiente (Requiere revisión con líder)'
                elif qual_stats['Mejoró'] > 0 or qual_stats['Mantuvo'] >= 2:
                    if mc_status == 'Mejoró':
                        real_status = 'Avance Integral Demostrado'
                    elif mc_status == 'Mantuvo':
                        real_status = 'Aprendizaje Consolidado'
                    else:
                        real_status = 'Avance Conceptual Cualitativo (con ajuste en alternativas)'
                elif mc_status == 'Retrocedió' and qual_stats['Retrocedió'] > 0:
                    real_status = 'Dificultad Conceptual (Requiere acompañamiento)'
                elif mc_status == 'Mejoró':
                    real_status = 'Mejora en Alternativas'
                else:
                    real_status = 'Rendimiento Estable'
                
                p_obj = {
                    'name': name,
                    'hito': hito_name,
                    'role': d.get('role', 'Estudiante'),
                    'school': sch,
                    'gender': gender,
                    'pre_score': pre_score,
                    'post_score': post_score,
                    'delta': delta,
                    'status': mc_status,
                    'real_status': real_status,
                    'qual_stats': qual_stats,
                    'q_evo': q_evo,
                    'pre_mc': d['PRE']['mc'],
                    'post_mc': d['POST']['mc'],
                    'open_analysis': open_analysis
                }
                paired_list.append(p_obj)
                all_participants[name]['hitos'][hito_name] = p_obj
                
        N = len(paired_list)
        status_c = Counter(p['status'] for p in paired_list)
        avg_pre = sum(p['pre_score'] for p in paired_list) / N if N else 0
        avg_post = sum(p['post_score'] for p in paired_list) / N if N else 0
        avg_delta = sum(p['delta'] for p in paired_list) / N if N else 0
        
        g_list = [(p['post_score'] - p['pre_score']) / (5.0 - p['pre_score']) for p in paired_list if p['pre_score'] < 5]
        hake_g = sum(g_list) / len(g_list) if g_list else 0
        
        questions_metrics = []
        for i in range(5):
            pre_c = sum(p['pre_mc'][i] == key[i] for p in paired_list)
            post_c = sum(p['post_mc'][i] == key[i] for p in paired_list)
            trans_c = Counter(p['q_evo'][i] for p in paired_list)
            post_distractors = Counter(p['post_mc'][i].upper() for p in paired_list if p['post_mc'][i] != key[i])
            
            questions_metrics.append({
                'num': i + 1,
                'title': q_titles[i],
                'correct_key': key[i].upper(),
                'pre_correct': pre_c,
                'pre_pct': round(pre_c / N * 100, 1) if N else 0,
                'post_correct': post_c,
                'post_pct': round(post_c / N * 100, 1) if N else 0,
                'diff': post_c - pre_c,
                'transitions': dict(trans_c),
                'post_distractors': dict(post_distractors)
            })
            
        by_school = defaultdict(list)
        for p in paired_list:
            by_school[p['school']].append(p)
            
        school_rankings = []
        for sname, plist in by_school.items():
            s_n = len(plist)
            s_pre = sum(p['pre_score'] for p in plist) / s_n
            s_post = sum(p['post_score'] for p in plist) / s_n
            s_del = sum(p['delta'] for p in plist) / s_n
            s_st = Counter(p['status'] for p in plist)
            meta = schools.get(sname, {})
            school_rankings.append({
                'school': sname,
                'n': s_n,
                'pre_avg': round(s_pre, 2),
                'post_avg': round(s_post, 2),
                'delta_avg': round(s_del, 2),
                'status_counts': dict(s_st),
                'pct_mejora': round(s_st['Mejoró'] / s_n * 100, 1) if s_n else 0,
                'lider1': meta.get('lider1', ''),
                'email1': meta.get('email1', ''),
                'tel1': meta.get('tel1', ''),
                'lider2': meta.get('lider2', ''),
                'email2': meta.get('email2', ''),
                'tel2': meta.get('tel2', ''),
                'dependencia': meta.get('dependencia', 'Municipal'),
                'certificacion': meta.get('certificacion', 'No'),
                'has_data': True
            })
            
        # Manejo de colegios de la Línea Verde sin evaluaciones en este hito específico (ej. Colegio Bet-el en Hito 4)
        active_school_names = set(by_school.keys())
        for sname in LINEA_VERDE_SCHOOLS.get(hito_name, []):
            if sname not in active_school_names:
                smeta = schools.get(sname, {})
                school_rankings.append({
                    'school': sname,
                    'n': 0,
                    'pre_avg': None,
                    'post_avg': None,
                    'delta_avg': 0.0,
                    'status_counts': {'Mejoró': 0, 'Mantuvo': 0, 'Retrocedió': 0},
                    'pct_mejora': 0.0,
                    'lider1': smeta.get('lider1', ''),
                    'email1': smeta.get('email1', ''),
                    'tel1': smeta.get('tel1', ''),
                    'lider2': smeta.get('lider2', ''),
                    'email2': smeta.get('email2', ''),
                    'tel2': smeta.get('tel2', ''),
                    'dependencia': smeta.get('dependencia', 'Particular Subvencionado'),
                    'certificacion': smeta.get('certificacion', 'No'),
                    'has_data': False,
                    'note': 'Sin datos registrados en este hito'
                })
                    
        school_rankings.sort(key=lambda x: (-1 if not x.get('has_data', True) else 0, -x['delta_avg']))
        
        by_gender = defaultdict(list)
        for p in paired_list:
            by_gender[p['gender']].append(p)
            
        gender_metrics = {}
        for gname, plist in by_gender.items():
            g_n = len(plist)
            g_pre = sum(p['pre_score'] for p in plist) / g_n
            g_post = sum(p['post_score'] for p in plist) / g_n
            g_del = sum(p['delta'] for p in plist) / g_n
            g_st = Counter(p['status'] for p in plist)
            gender_metrics[gname] = {
                'n': g_n,
                'pre_avg': round(g_pre, 2),
                'post_avg': round(g_post, 2),
                'delta_avg': round(g_del, 2),
                'status_counts': dict(g_st),
                'pct_mejora': round(g_st['Mejoró'] / g_n * 100, 1) if g_n else 0
            }
            
        hitos_data[hito_name] = {
            'desafio': desafio_name,
            'n_participants': N,
            'avg_pre': round(avg_pre, 2),
            'avg_post': round(avg_post, 2),
            'avg_delta': round(avg_delta, 2),
            'hake_gain': round(hake_g, 3),
            'status_counts': dict(status_c),
            'questions': questions_metrics,
            'open_questions': open_q_titles,
            'schools': school_rankings,
            'gender': gender_metrics,
            'participants': paired_list
        }

    # 3. Resumen Global consolidado (Hito 3 y Hito 4)
    all_paired_glob = []
    for hname in ['Hito 3', 'Hito 4']:
        if hname in hitos_data:
            for p in hitos_data[hname]['participants']:
                all_paired_glob.append(p)

    N_glob = len(all_paired_glob)
    if N_glob > 0:
        avg_pre_glob = sum(p['pre_score'] for p in all_paired_glob) / N_glob
        avg_post_glob = sum(p['post_score'] for p in all_paired_glob) / N_glob
        avg_delta_glob = sum(p['delta'] for p in all_paired_glob) / N_glob
        status_c_glob = Counter(p['status'] for p in all_paired_glob)
        g_list_glob = [(p['post_score'] - p['pre_score']) / (5.0 - p['pre_score']) for p in all_paired_glob if p['pre_score'] < 5]
        hake_g_glob = sum(g_list_glob) / len(g_list_glob) if g_list_glob else 0

        # Colegios consolidados
        by_school_glob = defaultdict(list)
        for p in all_paired_glob:
            by_school_glob[p['school']].append(p)

        schools_glob = []
        for sname, plist in by_school_glob.items():
            s_n = len(plist)
            s_pre = sum(p['pre_score'] for p in plist) / s_n
            s_post = sum(p['post_score'] for p in plist) / s_n
            s_del = sum(p['delta'] for p in plist) / s_n
            s_st = Counter(p['status'] for p in plist)
            meta = schools.get(sname, {})
            schools_glob.append({
                'school': sname,
                'n': s_n,
                'pre_avg': round(s_pre, 2),
                'post_avg': round(s_post, 2),
                'delta_avg': round(s_del, 2),
                'status_counts': dict(s_st),
                'pct_mejora': round(s_st['Mejoró'] / s_n * 100, 1) if s_n else 0,
                'lider1': meta.get('lider1', ''),
                'email1': meta.get('email1', ''),
                'tel1': meta.get('tel1', ''),
                'lider2': meta.get('lider2', ''),
                'email2': meta.get('email2', ''),
                'tel2': meta.get('tel2', ''),
                'dependencia': meta.get('dependencia', 'Municipal'),
                'certificacion': meta.get('certificacion', 'No'),
                'has_data': True
            })
        schools_glob.sort(key=lambda x: -x['delta_avg'])

        # Género consolidado
        by_gender_glob = defaultdict(list)
        for p in all_paired_glob:
            by_gender_glob[p['gender']].append(p)

        gender_metrics_glob = {}
        for gname, plist in by_gender_glob.items():
            g_n = len(plist)
            g_pre = sum(p['pre_score'] for p in plist) / g_n
            g_post = sum(p['post_score'] for p in plist) / g_n
            g_del = sum(p['delta'] for p in plist) / g_n
            g_st = Counter(p['status'] for p in plist)
            gender_metrics_glob[gname] = {
                'n': g_n,
                'pre_avg': round(g_pre, 2),
                'post_avg': round(g_post, 2),
                'delta_avg': round(g_del, 2),
                'status_counts': dict(g_st),
                'pct_mejora': round(g_st['Mejoró'] / g_n * 100, 1) if g_n else 0
            }

        glob_questions = []
        if 'Hito 3' in hitos_data:
            for q in hitos_data['Hito 3']['questions']:
                qc = dict(q)
                qc['hito'] = 'Hito 3'
                qc['title'] = f"[Hito 3] {qc['title']}"
                glob_questions.append(qc)
        if 'Hito 4' in hitos_data:
            for q in hitos_data['Hito 4']['questions']:
                qc = dict(q)
                qc['hito'] = 'Hito 4'
                qc['title'] = f"[Hito 4] {qc['title']}"
                glob_questions.append(qc)

        glob_open = []
        if 'Hito 3' in hitos_data:
            glob_open.extend([f"[Hito 3] {t}" for t in hitos_data['Hito 3']['open_questions']])
        if 'Hito 4' in hitos_data:
            glob_open.extend([f"[Hito 4] {t}" for t in hitos_data['Hito 4']['open_questions']])

        hitos_data['Global'] = {
            'desafio': 'Resumen Global Línea Verde (Hito 3 y Hito 4)',
            'n_participants': N_glob,
            'avg_pre': round(avg_pre_glob, 2),
            'avg_post': round(avg_post_glob, 2),
            'avg_delta': round(avg_delta_glob, 2),
            'hake_gain': round(hake_g_glob, 3),
            'status_counts': dict(status_c_glob),
            'questions': glob_questions,
            'open_questions': glob_open,
            'schools': schools_glob,
            'gender': gender_metrics_glob,
            'participants': all_paired_glob
        }

    return {
        'metadata': {
            'programa': 'Programa de Acción Climática (PAC) 2026',
            'linea': 'Línea Verde',
            'archivo_fuente': os.path.basename(file_path),
            'total_participantes': len(all_participants)
        },
        'schools_directory': schools,
        'hitos': hitos_data,
        'all_participants': list(all_participants.values())
    }

if __name__ == '__main__':
    import json
    current_dir = os.path.dirname(os.path.abspath(__file__))
    excel_path = os.path.join(current_dir, 'ParticipantesPAC2026_live.xlsx')
    if not os.path.exists(excel_path):
        excel_path = os.path.join(current_dir, '..', 'datos_evaluacion', 'Participantes PAC 2026.xlsx')
    parsed = parse_excel_data(excel_path)
    out_file = os.path.join(current_dir, 'data.json')
    with open(out_file, 'w', encoding='utf-8') as f:
        json.dump(parsed, f, ensure_ascii=False, indent=2)
    print(f"data.json generado correctamente en: {out_file}")
    print("Hitos disponibles:", list(parsed['hitos'].keys()))
    print(f"Total evaluaciones globales: {parsed['hitos']['Global']['n_participants']}")

